"""Report generation: CSV, XLSX and PDF.

Reports always render the *stored* data (database records or saved plan
results). Nothing is recalculated here unless the caller explicitly asks for it
elsewhere. All user-controlled strings are sanitized before being embedded in
PDF markup.
"""

from __future__ import annotations

import csv
import io
import ipaddress
from datetime import datetime, timezone
from typing import Any
from xml.sax.saxutils import escape

from sqlalchemy.orm import Session

from app.core.errors import not_found
from app.models import AllocationPlan, Network
from app.services.subnet_calculator import parse_network, subnet_details

ALLOCATION_COLUMNS = [
    "sequence",
    "department",
    "vlan_id",
    "requested_hosts",
    "effective_hosts",
    "network",
    "prefix_length",
    "subnet_mask",
    "wildcard_mask",
    "network_address",
    "first_usable",
    "last_usable",
    "broadcast_address",
    "total_addresses",
    "usable_capacity",
    "unused_usable_capacity",
    "utilization_percent",
    "gateway",
    "reserved_addresses",
    "status",
]


def _fmt(value: Any) -> str:
    return "" if value is None else str(value)


# ---------------------------------------------------------------------------
# Data collection (from stored records, no recalculation)
# ---------------------------------------------------------------------------

def network_report_data(db: Session, network_id: int) -> dict[str, Any]:
    network = db.get(Network, network_id)
    if not network:
        raise not_found("Network", network_id)
    parent_net = parse_network(network.cidr)
    subnets = sorted(
        [s for s in network.subnets if s.status != "released"],
        key=lambda s: int(ipaddress.ip_network(s.cidr).network_address),
    )
    rows: list[dict[str, Any]] = []
    allocated_total = 0
    for idx, s in enumerate(subnets, start=1):
        net = parse_network(s.cidr)
        details = subnet_details(net)
        allocated_total += net.num_addresses
        rows.append(
            {
                "sequence": idx,
                "department": s.name,
                "vlan_id": s.vlan_id,
                "network": s.cidr,
                "prefix_length": net.prefixlen,
                "subnet_mask": details.get("subnet_mask", ""),
                "wildcard_mask": details.get("wildcard_mask", ""),
                "network_address": str(net.network_address),
                "first_usable": details.get("first_usable", ""),
                "last_usable": details.get("last_usable", ""),
                "broadcast_address": details.get("broadcast_address", ""),
                "total_addresses": net.num_addresses,
                "usable_capacity": details.get("usable_hosts"),
                "gateway": s.gateway or "",
                "status": s.status,
                "description": s.description or "",
                "ip_address_count": len(s.ip_addresses),
            }
        )
    total = parent_net.num_addresses
    return {
        "kind": "network",
        "title": f"Network Report — {network.name}",
        "network": {
            "id": network.id,
            "name": network.name,
            "cidr": network.cidr,
            "family": f"IPv{network.family}",
            "environment": network.environment,
            "site": network.site or "",
            "status": network.status,
            "description": network.description or "",
        },
        "allocations": rows,
        "summary": {
            "parent_total_addresses": total,
            "allocated_addresses": allocated_total,
            "remaining_addresses": total - allocated_total,
            "subnet_count": len(rows),
            "overall_utilization_percent": round(100.0 * allocated_total / total, 2) if total else 0.0,
        },
        "remaining_ranges": _remaining_from_subnets(parent_net, [parse_network(s.cidr) for s in subnets]),
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


def plan_report_data(db: Session, plan_id: int) -> dict[str, Any]:
    plan = db.get(AllocationPlan, plan_id)
    if not plan:
        raise not_found("AllocationPlan", plan_id)
    results = plan.results or {}
    return {
        "kind": "plan",
        "title": f"Allocation Plan — {plan.name}",
        "plan": {
            "id": plan.id,
            "name": plan.name,
            "base_network": plan.base_network,
            "family": f"IPv{plan.family}",
            "status": plan.status,
            "algorithm_version": plan.algorithm_version,
            "notes": plan.notes or "",
        },
        "allocations": [
            {k: a.get(k) for k in ALLOCATION_COLUMNS} for a in results.get("allocations", [])
        ],
        "summary": results.get("summary", {}),
        "remaining_ranges": results.get("remaining_ranges", []),
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


def _remaining_from_subnets(parent: Any, allocated: list[Any]) -> list[dict[str, Any]]:
    from app.services.subnet_calculator import remaining_ranges

    return remaining_ranges(parent, allocated)


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def render_csv(data: dict[str, Any]) -> str:
    """Standards-compliant CSV with a metadata preamble comment block and the
    allocation table. Deterministic row ordering."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow([f"# {data['title']}"])
    if data["kind"] == "network":
        net = data["network"]
        writer.writerow(["# Network", net["name"], net["cidr"], net["family"], net["environment"], net["site"]])
    else:
        plan = data["plan"]
        writer.writerow(["# Plan", plan["name"], plan["base_network"], plan["family"], plan["status"]])
    summary = data.get("summary", {})
    writer.writerow(
        ["# Generated", data["generated_at"],
         f"total={summary.get('parent_total_addresses', '')}",
         f"allocated={summary.get('allocated_addresses', '')}",
         f"remaining={summary.get('remaining_addresses', '')}",
         f"utilization={summary.get('overall_utilization_percent', '')}%"]
    )
    writer.writerow([])
    columns = list(ALLOCATION_COLUMNS)
    writer.writerow(columns)
    for row in data["allocations"]:
        writer.writerow([_fmt(row.get(c)) for c in columns])
    writer.writerow([])
    writer.writerow(["Remaining Ranges", "Start", "End", "Total Addresses", "Usable Hosts"])
    for r in data.get("remaining_ranges", []):
        writer.writerow(
            [_fmt(r.get("start")), _fmt(r.get("end")), _fmt(r.get("total_addresses")), _fmt(r.get("usable_hosts"))]
        )
    return buf.getvalue()


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def render_xlsx(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    title_font = Font(bold=True, size=14)

    def style_header(ws, row: int, col_count: int) -> None:
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def autosize(ws) -> None:
        for col_cells in ws.columns:
            length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 3, 40)

    # --- Summary sheet
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = data["title"]
    ws["A1"].font = title_font
    row = 3
    if data["kind"] == "network":
        net = data["network"]
        pairs = [
            ("Network name", net["name"]),
            ("CIDR", net["cidr"]),
            ("Address family", net["family"]),
            ("Environment", net["environment"]),
            ("Site", net["site"]),
            ("Status", net["status"]),
            ("Description", net["description"]),
        ]
    else:
        plan = data["plan"]
        pairs = [
            ("Plan name", plan["name"]),
            ("Base network", plan["base_network"]),
            ("Address family", plan["family"]),
            ("Plan status", plan["status"]),
            ("Algorithm version", plan["algorithm_version"]),
            ("Notes", plan["notes"]),
        ]
    summary = data.get("summary", {})
    pairs += [
        ("Parent total addresses", summary.get("parent_total_addresses")),
        ("Allocated addresses", summary.get("allocated_addresses")),
        ("Remaining addresses", summary.get("remaining_addresses")),
        ("Overall utilization %", summary.get("overall_utilization_percent")),
        ("Generated at", data["generated_at"]),
    ]
    for label, value in pairs:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=row, column=2, value=_fmt(value))
        cell.number_format = "@"  # keep everything (esp. IP/CIDR) as text
        row += 1
    warnings = summary.get("warnings") or []
    if warnings:
        row += 1
        ws.cell(row=row, column=1, value="Warnings").font = Font(bold=True)
        for w in warnings:
            row += 1
            ws.cell(row=row, column=2, value=_fmt(w)).number_format = "@"
    autosize(ws)

    # --- Allocation sheet
    ws = wb.create_sheet("VLSM Allocation")
    columns = ALLOCATION_COLUMNS
    ws.append(columns)
    style_header(ws, 1, len(columns))
    for alloc in data["allocations"]:
        ws.append([_fmt(alloc.get(c)) for c in columns])
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(columns) + 1):
            ws.cell(row=r, column=c).number_format = "@"
    ws.freeze_panes = "A2"
    autosize(ws)

    # --- Remaining ranges sheet
    ws = wb.create_sheet("Remaining Ranges")
    rem_cols = ["start", "end", "total_addresses", "usable_hosts", "cidrs"]
    ws.append(rem_cols)
    style_header(ws, 1, len(rem_cols))
    for r in data.get("remaining_ranges", []):
        ws.append([
            _fmt(r.get("start")),
            _fmt(r.get("end")),
            _fmt(r.get("total_addresses")),
            _fmt(r.get("usable_hosts")),
            ", ".join(r.get("cidrs", [])),
        ])
    for rr in range(2, ws.max_row + 1):
        for c in range(1, len(rem_cols) + 1):
            ws.cell(row=rr, column=c).number_format = "@"
    ws.freeze_panes = "A2"
    autosize(ws)

    # --- Network metadata sheet
    ws = wb.create_sheet("Network Metadata")
    ws.append(["Field", "Value"])
    style_header(ws, 1, 2)
    if data["kind"] == "network":
        meta_rows = [
            ("Network", data["network"]["cidr"]),
            ("Family", data["network"]["family"]),
            ("Environment", data["network"]["environment"]),
            ("Site", data["network"]["site"]),
            ("Subnet count", summary.get("subnet_count")),
        ]
    else:
        meta_rows = [
            ("Base network", data["plan"]["base_network"]),
            ("Family", data["plan"]["family"]),
            ("Algorithm version", data["plan"]["algorithm_version"]),
            ("Plan status", data["plan"]["status"]),
            ("Department count", len(data["allocations"])),
        ]
    for label, value in meta_rows:
        ws.append([label, _fmt(value)])
    for rr in range(2, ws.max_row + 1):
        ws.cell(row=rr, column=2).number_format = "@"
    autosize(ws)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def render_pdf(data: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()

    from reportlab.pdfgen import canvas as _rl_canvas

    class _PlainCanvas(_rl_canvas.Canvas):
        """Force uncompressed content streams so exports are greppable."""

        def __init__(self, *args, **kwargs):  # type: ignore[no-untyped-def]
            kwargs["pageCompression"] = 0
            super().__init__(*args, **kwargs)

    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=data["title"][:100],
        author="NetSpace IPAM",
        pageCompression=0,
        canvasmaker=_PlainCanvas,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Heading1"], alignment=TA_CENTER, fontSize=16, spaceAfter=6
    )
    h2_style = ParagraphStyle("ReportH2", parent=styles["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4)
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=7, leading=9, wordWrap="CJK")

    story: list[Any] = []
    story.append(Paragraph(escape(data["title"]), title_style))
    story.append(Spacer(1, 4))

    if data["kind"] == "network":
        net = data["network"]
        info = [
            ("Parent network", net["cidr"]),
            ("Network name", net["name"]),
            ("Address family", net["family"]),
            ("Environment", net["environment"]),
            ("Site", net["site"] or "-"),
            ("Allocation date", data["generated_at"][:19].replace("T", " ")),
        ]
    else:
        plan = data["plan"]
        info = [
            ("Base network", plan["base_network"]),
            ("Plan name", plan["name"]),
            ("Address family", plan["family"]),
            ("Plan status", plan["status"]),
            ("Algorithm version", plan["algorithm_version"]),
            ("Allocation date", data["generated_at"][:19].replace("T", " ")),
        ]
    info_table = Table(
        [[Paragraph(f"<b>{escape(k)}</b>", cell_style), Paragraph(escape(str(v)), cell_style)] for k, v in info],
        colWidths=[45 * mm, 190 * mm],
    )
    info_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(info_table)
    story.append(Spacer(1, 8))

    # --- allocation table (headers repeat across pages via repeatRows)
    story.append(Paragraph("VLSM Allocation Table", h2_style))
    pdf_columns = [
        ("sequence", "#", 8),
        ("department", "Department", 34),
        ("vlan_id", "VLAN", 10),
        ("requested_hosts", "Req. hosts", 14),
        ("effective_hosts", "Eff. hosts", 14),
        ("network", "Network/CIDR", 34),
        ("subnet_mask", "Subnet mask", 28),
        ("first_usable", "First usable", 28),
        ("last_usable", "Last usable", 28),
        ("broadcast_address", "Broadcast", 28),
        ("total_addresses", "Total", 12),
        ("usable_capacity", "Usable", 12),
        ("utilization_percent", "Util %", 10),
        ("status", "Status", 16),
    ]
    header = [Paragraph(f"<b>{escape(title)}</b>", cell_style) for _, title, _ in pdf_columns]
    rows = [header]
    for alloc in data["allocations"]:
        rows.append(
            [
                Paragraph(escape(_fmt(alloc.get(key))), cell_style)
                for key, _, _ in pdf_columns
            ]
        )
    col_widths = [w * mm for _, _, w in pdf_columns]
    table = Table(rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 8))

    # --- summary
    summary = data.get("summary", {})
    story.append(Paragraph("Network Utilization Summary", h2_style))
    summary_rows = [
        ["Total parent addresses", _fmt(summary.get("parent_total_addresses"))],
        ["Allocated addresses", _fmt(summary.get("allocated_addresses"))],
        ["Remaining addresses", _fmt(summary.get("remaining_addresses"))],
        ["Overall utilization", f"{_fmt(summary.get('overall_utilization_percent'))}%"],
    ]
    if data["kind"] == "plan":
        summary_rows += [
            ["Total requested hosts", _fmt(summary.get("total_requested_hosts"))],
            ["Total effective hosts", _fmt(summary.get("total_effective_hosts"))],
        ]
    sum_table = Table(
        [[Paragraph(f"<b>{escape(k)}</b>", cell_style), Paragraph(escape(v), cell_style)] for k, v in summary_rows],
        colWidths=[60 * mm, 60 * mm],
    )
    sum_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(sum_table)

    # --- remaining address space
    remaining = data.get("remaining_ranges", [])
    if remaining:
        story.append(Paragraph("Remaining Address Space", h2_style))
        rem_rows = [["Start", "End", "Total addresses", "Usable hosts"]]
        for r in remaining:
            rem_rows.append(
                [
                    _fmt(r.get("start")),
                    _fmt(r.get("end")),
                    _fmt(r.get("total_addresses")),
                    _fmt(r.get("usable_hosts")),
                ]
            )
        rem_table = Table(
            [[Paragraph(escape(str(v)), cell_style) for v in row] for row in rem_rows],
            colWidths=[40 * mm, 40 * mm, 35 * mm, 35 * mm],
            repeatRows=1,
        )
        rem_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4C6E8F")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.append(rem_table)

    # --- warnings + notes
    warnings = summary.get("warnings") or []
    if warnings:
        story.append(Paragraph("Allocation Warnings", h2_style))
        for w in warnings:
            story.append(Paragraph(f"• {escape(str(w))}", cell_style))
    if data["kind"] == "plan":
        notes = data["plan"].get("notes") or ""
        if notes:
            story.append(Paragraph("Plan Notes", h2_style))
            story.append(Paragraph(escape(notes), cell_style))
    story.append(Spacer(1, 8))
    story.append(
        Paragraph(
            "<i>Algorithm: largest-first aligned first-fit VLSM. IPv4 usable capacity excludes "
            "network and broadcast addresses (/31 point-to-point and /32 host routes follow RFC 3021 "
            "semantics). IPv6 subnets have no broadcast address; capacity is reported as the total "
            "address count.</i>",
            cell_style,
        )
    )

    def _footer(canvas, doc_):  # type: ignore[no-untyped-def]
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(12 * mm, 6 * mm, "NetSpace IPAM — generated report")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 6 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()
